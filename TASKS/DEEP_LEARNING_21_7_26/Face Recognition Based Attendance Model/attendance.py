"""
attendance.py
-------------
STEP 4 + 5 of the pipeline:
  - Opens the webcam
  - Detects faces frame-by-frame using OpenCV / face_recognition
  - Generates a 128-d vector for each detected face
  - Queries ChromaDB for the nearest matching student vector
  - Marks attendance (1) for recognized students in an Excel sheet
  - At the end of the session, marks everyone NOT seen as absent (0)

Run:
    python attendance.py

Controls:
    Press 'q' to end the session and finalize the Excel sheet.
"""

import cv2
import pickle
import os
import datetime
import face_recognition
import numpy as np
from openpyxl import Workbook, load_workbook

from vector_store import get_collection, ENCODINGS_FILE

ATTENDANCE_FILE = "attendance_records/attendance.xlsx"
MATCH_THRESHOLD = 0.6   # lower = stricter match (face_recognition's typical default)
RESIZE_SCALE = 0.25     # shrink frame for faster detection


def load_all_student_names():
    """Reads the pickle file just to get the full roster (for marking absentees)."""
    if not os.path.exists(ENCODINGS_FILE):
        return []
    with open(ENCODINGS_FILE, "rb") as f:
        data = pickle.load(f)
    return [(name, info["roll_no"]) for name, info in data.items()]


def init_attendance_sheet():
    os.makedirs(os.path.dirname(ATTENDANCE_FILE), exist_ok=True)
    if os.path.exists(ATTENDANCE_FILE):
        wb = load_workbook(ATTENDANCE_FILE)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.append(["Name", "Roll No", "Date", "Time", "Status"])
        wb.save(ATTENDANCE_FILE)
    return wb


def mark_attendance(name, roll_no, status):
    wb = init_attendance_sheet()
    ws = wb["Attendance"]

    today = datetime.date.today().isoformat()
    now = datetime.datetime.now().strftime("%H:%M:%S")

    # avoid duplicate row for same person/date
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value == name and row[2].value == today:
            return  # already logged today

    ws.append([name, roll_no, today, now, status])
    wb.save(ATTENDANCE_FILE)


def finalize_absentees(seen_names):
    """Any enrolled student not seen during the session gets marked absent (0)."""
    all_students = load_all_student_names()
    for name, roll_no in all_students:
        if name not in seen_names:
            mark_attendance(name, roll_no, 0)


def run_attendance_session():
    collection = get_collection()

    cap = cv2.VideoCapture(0)
    if not cap.isOpened():
        raise RuntimeError("Could not open webcam. Check camera permissions/index.")

    seen_today = set()
    print("[INFO] Attendance session started. Press 'q' to finish and save the sheet.")

    while True:
        ret, frame = cap.read()
        if not ret:
            break

        small_frame = cv2.resize(frame, (0, 0), fx=RESIZE_SCALE, fy=RESIZE_SCALE)
        rgb_small_frame = cv2.cvtColor(small_frame, cv2.COLOR_BGR2RGB)

        face_locations = face_recognition.face_locations(rgb_small_frame, model="hog")
        face_encodings = face_recognition.face_encodings(rgb_small_frame, face_locations)

        for (top, right, bottom, left), face_encoding in zip(face_locations, face_encodings):
            result = collection.query(
                query_embeddings=[face_encoding.tolist()],
                n_results=1
            )

            name_label = "Unknown"
            color = (0, 0, 255)  # red for unknown

            if result["distances"] and result["distances"][0]:
                distance = result["distances"][0][0]
                # chromadb l2 space returns squared L2; face_recognition typically
                # compares on plain L2, so we take sqrt for a comparable threshold
                l2_distance = distance ** 0.5

                if l2_distance < MATCH_THRESHOLD:
                    meta = result["metadatas"][0][0]
                    name_label = f'{meta["name"]} ({meta["roll_no"]})'
                    color = (0, 255, 0)  # green for recognized

                    if meta["name"] not in seen_today:
                        seen_today.add(meta["name"])
                        mark_attendance(meta["name"], meta["roll_no"], 1)
                        print(f'[ATTENDANCE] Marked present: {meta["name"]} ({meta["roll_no"]})')

            # scale coordinates back up to original frame size
            top = int(top / RESIZE_SCALE)
            right = int(right / RESIZE_SCALE)
            bottom = int(bottom / RESIZE_SCALE)
            left = int(left / RESIZE_SCALE)

            cv2.rectangle(frame, (left, top), (right, bottom), color, 2)
            cv2.putText(frame, name_label, (left, top - 10),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.6, color, 2)

        cv2.imshow("Attendance - press 'q' to finish", frame)
        if cv2.waitKey(1) & 0xFF == ord('q'):
            break

    cap.release()
    cv2.destroyAllWindows()

    finalize_absentees(seen_today)
    print(f"[SUCCESS] Session finished. Attendance saved to {ATTENDANCE_FILE}")
    print(f"[INFO] Present today: {sorted(seen_today) if seen_today else 'None'}")


if __name__ == "__main__":
    run_attendance_session()
